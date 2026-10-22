"""
Catálogo de referência (SUPER_ADMIN): the platform's products by sector (see app.models.reference_product) - typed
one by one or imported from an Excel sheet (one row per product, its sale units on the same row). A product is known by
its barcode, or by its name when it has none: importing the sheet again updates the products already there.
"""
import io
import uuid
from decimal import Decimal, InvalidOperation

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from sqlalchemy import func, or_, select
from sqlalchemy.ext.asyncio import AsyncSession

from app.core.capabilities import SECTORS
from app.models.reference_product import ReferenceProduct, ReferenceProductUnit
from app.models.unit_of_measure_catalog import UnitOfMeasureCatalog
from app.models.vat_code import VatCode

TAX_CATEGORIES = {"NOR": "Taxa normal", "RED": "Taxa reduzida", "ISE": "Isento"}
MAX_SALE_UNITS = 3  # on a sheet row; the screen takes as many as needed
GENERAL = "GERAL"  # the list's filter for the products without sector (offered to every sector)


class ReferenceCatalogError(Exception):
    """Raised when a reference product is refused (its name, barcode, VAT nature, sectors, units)."""


def catalog_sectors() -> list[dict]:
    return [{"code": s.code, "name": s.name} for s in SECTORS if s.ready]


async def catalog_meta(db: AsyncSession) -> dict:
    """What the screen chooses from: the sectors (with their product counts), the units, the exemption motives."""
    counts = dict((await db.execute(
        select(func.unnest(ReferenceProduct.sectors), func.count()).group_by(func.unnest(ReferenceProduct.sectors))
    )).all())
    units = (await db.execute(select(UnitOfMeasureCatalog).where(UnitOfMeasureCatalog.is_active.is_(True))
                              .order_by(UnitOfMeasureCatalog.code))).scalars().all()
    # a product's own motives only: a regime's motive (M00, M04) is the company's, set by its regime when it takes it
    from app.models.fiscal_regime import FiscalRegime
    regime_motives = select(FiscalRegime.required_exemption_id).where(FiscalRegime.required_exemption_id.is_not(None))
    motives = (await db.execute(select(VatCode).where(VatCode.is_active.is_(True), VatCode.id.not_in(regime_motives))
                                .order_by(VatCode.code))).scalars().all()
    total = (await db.execute(select(func.count()).select_from(ReferenceProduct))).scalar_one()
    general = (await db.execute(select(func.count()).select_from(ReferenceProduct)
                                .where(func.cardinality(ReferenceProduct.sectors) == 0))).scalar_one()
    return {
        "sectors": [{**s, "count": counts.get(s["code"], 0)} for s in catalog_sectors()],
        "total": total,
        "general": general,  # products without sector: offered to every sector
        "units": [{"id": u.id, "code": u.code, "name": u.name} for u in units],
        "exemptions": [{"id": m.id, "code": m.code, "name": m.name} for m in motives],
        "tax_categories": [{"code": k, "name": v} for k, v in TAX_CATEGORIES.items()],
    }


def product_view(p: ReferenceProduct) -> dict:
    return {
        "id": p.id, "barcode": p.barcode, "name": p.name, "category": p.category,
        "unit_of_measure_id": p.unit_of_measure_id, "tax_category": p.tax_category, "exemption_id": p.exemption_id,
        "price": float(p.price) if p.price is not None else None, "sectors": p.sectors, "is_active": p.is_active,
        "units": [{"id": u.id, "unit_of_measure_id": u.unit_of_measure_id, "factor": float(u.factor), "barcode": u.barcode,
                   "price": float(u.price) if u.price is not None else None} for u in p.units],
        "updated_at": p.updated_at,
    }


async def list_reference_products(db: AsyncSession, q: str = "", sector: str | None = None, limit: int = 100,
                                  offset: int = 0) -> dict:
    """A page of the catalog: searched by name, barcode (of the product or of a sale unit) or category."""
    query = select(ReferenceProduct)
    term = (q or "").strip()
    if term:
        like = f"%{term.lower()}%"
        query = query.where(or_(
            func.lower(ReferenceProduct.name).like(like), ReferenceProduct.barcode == term,
            func.lower(ReferenceProduct.category).like(like),
            ReferenceProduct.id.in_(select(ReferenceProductUnit.reference_product_id).where(ReferenceProductUnit.barcode == term)),
        ))
    if sector == GENERAL:
        query = query.where(func.cardinality(ReferenceProduct.sectors) == 0)
    elif sector:
        query = query.where(ReferenceProduct.sectors.any(sector))
    total = (await db.execute(select(func.count()).select_from(query.subquery()))).scalar_one()
    rows = (await db.execute(query.order_by(func.lower(ReferenceProduct.name)).limit(limit).offset(offset))).scalars().all()
    return {"total": total, "items": [product_view(p) for p in rows]}


def _money(value) -> Decimal | None:
    if value is None or value == "":
        return None
    try:
        amount = Decimal(str(value).replace(",", "."))
    except InvalidOperation:
        raise ReferenceCatalogError(f"Preço inválido: {value}")
    if amount < 0:
        raise ReferenceCatalogError("O preço não pode ser negativo")
    return amount


async def save_reference_product(db: AsyncSession, data: dict, product_id: uuid.UUID | None = None,
                                 commit: bool = True) -> ReferenceProduct:
    """Creates or changes a reference product, its sale units replaced by the ones given. Changes only the catalog:
    the companies that took the product keep their own copy."""
    name = (data.get("name") or "").strip()
    if not name:
        raise ReferenceCatalogError("Indique o nome do produto")
    barcode = (data.get("barcode") or "").strip() or None
    # the sectors whose new companies receive it - none: every sector (water, soap...)
    sectors = sorted({s for s in (data.get("sectors") or []) if s})
    known = {s["code"] for s in catalog_sectors()}
    if set(sectors) - known:
        raise ReferenceCatalogError(f"Setor desconhecido: {sorted(set(sectors) - known)[0]}")
    tax = (data.get("tax_category") or "").upper()
    if tax not in TAX_CATEGORIES:
        raise ReferenceCatalogError("Escolha a natureza do IVA: NOR (normal), RED (reduzida) ou ISE (isento)")
    exemption_id = data.get("exemption_id") or None
    if tax == "ISE" and exemption_id is None:
        raise ReferenceCatalogError("Um produto isento precisa do motivo de isenção")
    if tax != "ISE":
        exemption_id = None
    elif await db.get(VatCode, exemption_id) is None:
        raise ReferenceCatalogError("Motivo de isenção desconhecido")
    unit_id = data.get("unit_of_measure_id")
    if unit_id is None or await db.get(UnitOfMeasureCatalog, unit_id) is None:
        raise ReferenceCatalogError("Escolha a unidade de base do produto")

    same = select(ReferenceProduct.id).where(ReferenceProduct.id != product_id) if product_id else select(ReferenceProduct.id)
    if barcode and (await db.execute(same.where(ReferenceProduct.barcode == barcode))).first():
        raise ReferenceCatalogError(f"Já existe um produto com o código de barras {barcode}")
    if not barcode and (await db.execute(same.where(ReferenceProduct.barcode.is_(None),
                                                    func.lower(ReferenceProduct.name) == name.lower()))).first():
        raise ReferenceCatalogError(f"Já existe um produto sem código de barras chamado «{name}»")

    async def barcode_taken(code: str) -> bool:
        """Another product of the catalog reads this barcode (as its own, or as one of its sale units')."""
        others = select(ReferenceProduct.id).where(ReferenceProduct.barcode == code)
        other_units = select(ReferenceProductUnit.id).where(ReferenceProductUnit.barcode == code)
        if product_id:
            others = others.where(ReferenceProduct.id != product_id)
            other_units = other_units.where(ReferenceProductUnit.reference_product_id != product_id)
        return bool((await db.execute(others)).first() or (await db.execute(other_units)).first())

    if barcode and (await db.execute(select(ReferenceProductUnit.id).where(
            ReferenceProductUnit.barcode == barcode,
            *( [ReferenceProductUnit.reference_product_id != product_id] if product_id else [])))).first():
        raise ReferenceCatalogError(f"O código de barras {barcode} já é de uma unidade de venda de outro produto")

    units, seen, unit_barcodes = [], set(), set()
    for raw in data.get("units") or []:
        unit = raw.get("unit_of_measure_id")
        if not unit:
            continue
        if unit == unit_id or str(unit) == str(unit_id):
            raise ReferenceCatalogError("Uma unidade de venda não pode ser a unidade de base")
        if str(unit) in seen:
            raise ReferenceCatalogError("A mesma unidade de venda aparece duas vezes")
        seen.add(str(unit))
        if await db.get(UnitOfMeasureCatalog, unit) is None:
            raise ReferenceCatalogError("Unidade de venda desconhecida")
        try:
            factor = Decimal(str(raw.get("factor")).replace(",", "."))
        except InvalidOperation:
            factor = Decimal(0)
        if factor <= 0:
            raise ReferenceCatalogError("Cada unidade de venda precisa da quantidade que contém (ex.: 6 unidades)")
        unit_barcode = (raw.get("barcode") or "").strip() or None
        if unit_barcode:  # one barcode reads one thing: never twice, never the product's own, never another product's
            if unit_barcode == barcode or unit_barcode in unit_barcodes:
                raise ReferenceCatalogError(f"O código de barras {unit_barcode} está repetido - cada unidade tem o seu")
            if await barcode_taken(unit_barcode):
                raise ReferenceCatalogError(f"O código de barras {unit_barcode} já é de outro produto do catálogo")
            unit_barcodes.add(unit_barcode)
        units.append(ReferenceProductUnit(unit_of_measure_id=unit, factor=factor, barcode=unit_barcode,
                                          price=_money(raw.get("price"))))

    if product_id:
        product = await db.get(ReferenceProduct, product_id)
        if product is None:
            raise ReferenceCatalogError("Produto não encontrado")
    else:
        product = ReferenceProduct()
        db.add(product)
    product.name, product.barcode, product.sectors = name, barcode, sectors
    product.category = (data.get("category") or "").strip() or None
    product.unit_of_measure_id, product.tax_category, product.exemption_id = unit_id, tax, exemption_id
    product.price = _money(data.get("price"))
    if "is_active" in data:
        product.is_active = bool(data["is_active"])
    if product_id and product.units:  # the old sale units out first: the same unit may come back
        product.units.clear()
        await db.flush()
    product.units = units
    await db.flush()
    if commit:
        await db.commit()
        await db.refresh(product)
    return product


async def delete_reference_product(db: AsyncSession, product_id: uuid.UUID) -> None:
    """Out of the catalog - the companies that took it keep their own copy."""
    product = await db.get(ReferenceProduct, product_id)
    if product is None:
        raise ReferenceCatalogError("Produto não encontrado")
    await db.delete(product)
    await db.commit()


# ------------------------------------------------------------------------------------------------------- Excel sheet

HEADERS = ["Código de barras", "Nome *", "Categoria", "Unidade *", "IVA *", "Motivo de isenção", "Preço", "Setores (vazio = todos)"]
for _n in range(1, MAX_SALE_UNITS + 1):
    HEADERS += [f"Unidade venda {_n}", f"Quantidade {_n}", f"Código barras {_n}", f"Preço {_n}"]


async def catalog_template(db: AsyncSession) -> bytes:
    """The sheet to fill: the columns, an example row, and the codes to use (units, motives, sectors)."""
    meta = await catalog_meta(db)
    book = Workbook()
    sheet = book.active
    sheet.title = "Produtos"
    sheet.append(HEADERS)
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="3D7FE6")
    sheet.append(["5601234567890", "Ovos frescos", "Ovos e lacticínios", "UN", "ISE", "M10", "", "COMERCIO",
                  "CX", 6, "5601234567891", "", "CT", 30, "", ""])
    for column, width in zip("ABCDEFGH", (18, 32, 22, 11, 7, 18, 10, 22)):
        sheet.column_dimensions[column].width = width
    codes = book.create_sheet("Códigos")
    codes.append(["Unidades (código)", "", "IVA", "", "Motivos de isenção", "", "Setores (código)", ""])
    rows = max(len(meta["units"]), len(meta["exemptions"]), len(meta["sectors"]), 3)
    taxes = list(TAX_CATEGORIES.items())
    for i in range(rows):
        u = meta["units"][i] if i < len(meta["units"]) else None
        m = meta["exemptions"][i] if i < len(meta["exemptions"]) else None
        s = meta["sectors"][i] if i < len(meta["sectors"]) else None
        t = taxes[i] if i < len(taxes) else None
        codes.append([u["code"] if u else "", u["name"] if u else "", t[0] if t else "", t[1] if t else "",
                      m["code"] if m else "", m["name"] if m else "", s["code"] if s else "", s["name"] if s else ""])
    out = io.BytesIO()
    book.save(out)
    return out.getvalue()


def _text(value) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)  # a barcode read as a number
    return str(value).strip()


async def import_catalog(db: AsyncSession, content: bytes) -> dict:
    """The sheet's rows into the catalog: a product already there (same barcode, else same name) is updated, a new one
    created. Each row stands alone: a refused row is reported with its reason, the others go in."""
    try:
        sheet = load_workbook(io.BytesIO(content), read_only=True, data_only=True).worksheets[0]
    except Exception:  # noqa: BLE001 - any unreadable file
        raise ReferenceCatalogError("Ficheiro Excel inválido - use o modelo («Descarregar modelo»)")
    units = {u.code.upper(): u.id for u in (await db.execute(select(UnitOfMeasureCatalog))).scalars().all()}
    motives = {m.code.upper(): m.id for m in (await db.execute(select(VatCode))).scalars().all()}
    sectors = {}
    for s in catalog_sectors():
        sectors[s["code"].upper()] = s["code"]
        sectors[s["name"].upper()] = s["code"]
    created = updated = 0
    errors = []
    for number, row in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
        cells = [_text(v) for v in row] + [""] * len(HEADERS)
        if not any(cells[:8]):
            continue
        try:
            def unit_of(code, label):
                if code.upper() not in units:
                    raise ReferenceCatalogError(f"{label} desconhecida: {code}")
                return units[code.upper()]
            barcode, name, category, unit, tax, motive, price, sector_text = cells[:8]
            sale_units = []
            for n in range(MAX_SALE_UNITS):
                code, factor, unit_barcode, unit_price = cells[8 + 4 * n: 12 + 4 * n]
                if code:
                    sale_units.append({"unit_of_measure_id": unit_of(code, "Unidade de venda"), "factor": factor or 0,
                                       "barcode": unit_barcode, "price": unit_price})
            wanted = [part.strip().upper() for part in sector_text.replace(";", ",").split(",") if part.strip()]
            unknown = [w for w in wanted if w not in sectors]
            if unknown:
                raise ReferenceCatalogError(f"Setor desconhecido: {unknown[0]}")
            if motive and motive.upper() not in motives:
                raise ReferenceCatalogError(f"Motivo de isenção desconhecido: {motive}")
            data = {"barcode": barcode, "name": name, "category": category,
                    "unit_of_measure_id": unit_of(unit, "Unidade") if unit else None, "tax_category": tax,
                    "exemption_id": motives.get(motive.upper()) if motive else None, "price": price,
                    "sectors": [sectors[w] for w in wanted], "units": sale_units}
            existing = (await db.execute(select(ReferenceProduct.id).where(
                ReferenceProduct.barcode == barcode if barcode else
                (ReferenceProduct.barcode.is_(None)) & (func.lower(ReferenceProduct.name) == name.lower())))).scalar_one_or_none()
            async with db.begin_nested():
                await save_reference_product(db, data, existing, commit=False)
            if existing:
                updated += 1
            else:
                created += 1
        except ReferenceCatalogError as e:
            errors.append({"row": number, "message": str(e)})
    await db.commit()
    return {"created": created, "updated": updated, "errors": errors}


# ------------------------------------------------------------------------------------- a new company takes the catalog

def _for_sectors(sectors: list[str]):
    """The active products a company of these sectors receives: theirs, and those of every sector (no sector)."""
    return select(ReferenceProduct).where(
        ReferenceProduct.is_active.is_(True),
        or_(func.cardinality(ReferenceProduct.sectors) == 0, ReferenceProduct.sectors.overlap(sectors or [""])),
    )


async def sectors_of_modules(db: AsyncSession, module_ids: list) -> list[str]:
    from app.models.module import Module
    if not module_ids:
        return []
    return [m.code for m in (await db.execute(select(Module).where(Module.id.in_(module_ids)))).scalars().all() if m.code]


async def catalog_offer(db: AsyncSession, module_ids: list) -> dict:
    """«Nova empresa» > Base: what the reference catalog gives a company of these sectors."""
    sectors = await sectors_of_modules(db, module_ids)
    rows = (await db.execute(_for_sectors(sectors))).scalars().all()
    names = {s["code"]: s["name"] for s in catalog_sectors()}
    return {"sectors": [names.get(s, s) for s in sectors], "products": len(rows),
            "without_price": sum(1 for p in rows if p.price is None)}


async def take_reference_catalog(db: AsyncSession, company) -> dict:
    """The new company receives the reference catalog of its sectors (see take_reference_products) - no commit."""
    from app.models.company_module import CompanyModule
    module_ids = (await db.execute(select(CompanyModule.module_id).where(CompanyModule.company_id == company.id))).scalars().all()
    sectors = await sectors_of_modules(db, list(module_ids))
    refs = (await db.execute(_for_sectors(sectors).order_by(ReferenceProduct.name))).scalars().all()
    counts, _ = await take_reference_products(db, company, refs)
    return counts


async def company_barcodes(db: AsyncSession, company_id) -> tuple[set, set]:
    """The barcodes a company already reads (its products' and their packs'), and its products' names (lowercase)."""
    from app.models.product import Product
    from app.models.product_sale_unit import ProductSaleUnit
    rows = (await db.execute(select(Product.barcode, Product.name).where(Product.company_id == company_id))).all()
    packs = (await db.execute(select(ProductSaleUnit.barcode).where(
        ProductSaleUnit.company_id == company_id, ProductSaleUnit.barcode.is_not(None)))).scalars().all()
    return {r.barcode for r in rows if r.barcode} | set(packs), {r.name.lower() for r in rows}


def already_owned(ref: ReferenceProduct, barcodes: set, names: set) -> bool:
    """The company has it: one of its barcodes (the product's or a pack's) is read already, or - without barcode - a
    product of that name."""
    codes = {ref.barcode, *(u.barcode for u in ref.units)} - {None}
    return bool(codes & barcodes) if codes else ref.name.lower() in names


async def take_reference_products(db: AsyncSession, company, refs) -> tuple[dict, list]:
    """These reference products become the company's own articles (no commit):
    each product becomes its own article - stock 0, its category, its sale units - with THE company's VAT:
    - the product's nature (NOR / RED / ISE) gives the company's rate of that category, when its regime allows it;
    - a regime imposing a motive (M00, M04) gives its exempt rate and that motive, unless the product's own motive is
      one the regime keeps;
    - otherwise the product is taken with the rate of its category, inactive: the reclassification banner shows it
      and it is never sold before (vat_rule_service).
    Without price: price 0, off the Caixa and never sold until the gestor sets it (awaiting_price).
    A product the company already has (one of its barcodes, else its name) is skipped. Returns the counts and the
    new articles' ids."""
    from app.models.fiscal_regime import FiscalRegime
    from app.models.legal_vat_rate import LegalVatRate
    from app.models.product import Product, ProductType
    from app.models.product_category import ProductCategory
    from app.models.product_sale_unit import ProductSaleUnit
    from app.models.vat import VAT

    regime = await db.get(FiscalRegime, company.fiscal_regime_id) if company.fiscal_regime_id else None
    rates = {}
    for vat in (await db.execute(select(VAT).where(VAT.company_id == company.id))).scalars().all():
        if vat.tax_category not in rates or vat.is_active:
            rates[vat.tax_category] = vat

    async def rate_of(category: str):
        """The company's rate of this category - an inactive one added when it has none (the product to reclassify)."""
        if category not in rates:
            legal = (await db.execute(select(LegalVatRate).where(LegalVatRate.tax_category == category))).scalar_one_or_none()
            vat = VAT(company_id=company.id, tax_category=category, is_active=False,
                      name=legal.name if legal else TAX_CATEGORIES.get(category, category), rate=legal.rate if legal else 0)
            db.add(vat)
            await db.flush()
            rates[category] = vat
        return rates[category]

    def allowed(category: str) -> bool:
        vat = rates.get(category)
        return vat is not None and vat.is_active and (regime is None or getattr(regime, "allows_" + category.lower(), False))

    async def vat_of(product: ReferenceProduct) -> tuple:
        """(rate, motive, to review) - THE rule of vat_rule_service, applied to a whole catalog."""
        imposed = regime.required_exemption_id if regime is not None else None
        kept = set(regime.kept_exemption_ids or []) if regime is not None else set()
        if allowed(product.tax_category):
            if product.tax_category != "ISE":
                return rates[product.tax_category].id, None, False
            motive = product.exemption_id if (imposed is None or product.exemption_id in kept) else imposed
            return rates["ISE"].id, motive, False
        if imposed is not None and allowed("ISE"):  # Simplificado, Exclusão: exempt with the regime's motive
            return rates["ISE"].id, product.exemption_id if product.exemption_id in kept else imposed, False
        return (await rate_of(product.tax_category)).id, product.exemption_id, True

    codes = set((await db.execute(select(Product.code).where(Product.company_id == company.id))).scalars().all())
    barcodes, names = await company_barcodes(db, company.id)
    categories = {c.name.lower(): c.id for c in (await db.execute(
        select(ProductCategory).where(ProductCategory.company_id == company.id))).scalars().all()}

    counts = {"artigos": 0, "sem_preco": 0, "a_rever": 0, "unidades_de_venda": 0, "categorias": 0}
    serial = 0
    created = []
    for ref in refs:
        if already_owned(ref, barcodes, names):
            continue
        category_id = None
        if ref.category:
            category_id = categories.get(ref.category.lower())
            if category_id is None:
                category = ProductCategory(company_id=company.id, name=ref.category)
                db.add(category)
                await db.flush()
                category_id = categories[ref.category.lower()] = category.id
                counts["categorias"] += 1
        code = ref.barcode if ref.barcode and ref.barcode not in codes and len(ref.barcode) <= 50 else None
        while code is None or code in codes:
            serial += 1
            code = f"REF{serial:05d}"
        codes.add(code)
        vat_id, motive, review = await vat_of(ref)
        no_price = ref.price is None
        product = Product(company_id=company.id, code=code, name=ref.name, barcode=ref.barcode, category_id=category_id,
                          unit_of_measure_id=ref.unit_of_measure_id, vat_id=vat_id, exemption_reason_id=motive,
                          price=ref.price if not no_price else 0, min_stock_threshold=0, product_type=ProductType.BEM,
                          awaiting_price=no_price, not_available_pos=no_price)
        db.add(product)
        await db.flush()
        for unit in ref.units:
            unit_price = unit.price if unit.price is not None else (Decimal(ref.price) * unit.factor if ref.price is not None else 0)
            db.add(ProductSaleUnit(company_id=company.id, product_id=product.id, unit_of_measure_id=unit.unit_of_measure_id,
                                   factor=unit.factor, price=unit_price, barcode=unit.barcode))
            counts["unidades_de_venda"] += 1
        created.append(product.id)
        counts["artigos"] += 1
        counts["sem_preco"] += no_price
        counts["a_rever"] += review
        barcodes.update({ref.barcode, *(u.barcode for u in ref.units)} - {None})
        names.add(ref.name.lower())
    await db.flush()
    return counts, created


# ------------------------------------------------------------------------- an existing company picks from the catalog

def company_item_view(ref: ReferenceProduct, owned: bool) -> dict:
    return {**product_view(ref), "owned": owned}


async def _company(db: AsyncSession, company_id):
    from app.models.company import Company
    return await db.get(Company, company_id)


async def company_catalog(db: AsyncSession, company_id, q: str = "", sector: str | None = None, mine: bool = True,
                          limit: int = 100, offset: int = 0) -> dict:
    """Artigos > «Importar do catálogo»: the active reference products - by default those of the company's sectors and
    of every sector («mine») - each marked when the company has it already."""
    from app.models.company_module import CompanyModule
    module_ids = (await db.execute(select(CompanyModule.module_id).where(CompanyModule.company_id == company_id))).scalars().all()
    sectors = await sectors_of_modules(db, list(module_ids))
    query = _for_sectors(sectors) if mine and not sector else select(ReferenceProduct).where(ReferenceProduct.is_active.is_(True))
    term = (q or "").strip()
    if term:
        like = f"%{term.lower()}%"
        query = query.where(or_(
            func.lower(ReferenceProduct.name).like(like), ReferenceProduct.barcode == term,
            func.lower(ReferenceProduct.category).like(like),
            ReferenceProduct.id.in_(select(ReferenceProductUnit.reference_product_id).where(ReferenceProductUnit.barcode == term)),
        ))
    if sector == GENERAL:
        query = query.where(func.cardinality(ReferenceProduct.sectors) == 0)
    elif sector:
        query = query.where(ReferenceProduct.sectors.any(sector))
    total = (await db.execute(select(func.count()).select_from(query.subquery()))).scalar_one()
    rows = (await db.execute(query.order_by(func.lower(ReferenceProduct.name)).limit(limit).offset(offset))).scalars().all()
    barcodes, names = await company_barcodes(db, company_id)
    names_of = {s["code"]: s["name"] for s in catalog_sectors()}
    return {"total": total, "my_sectors": [{"code": s, "name": names_of.get(s, s)} for s in sectors],
            "sectors": catalog_sectors(),
            "items": [company_item_view(r, already_owned(r, barcodes, names)) for r in rows]}


async def lookup_barcode(db: AsyncSession, company_id, barcode: str) -> dict | None:
    """Artigos > «Novo produto»: a barcode typed or scanned that the reference catalog knows (the product's own, or one
    of its packs')."""
    code = (barcode or "").strip()
    if not code:
        return None
    ref = (await db.execute(select(ReferenceProduct).where(ReferenceProduct.is_active.is_(True), or_(
        ReferenceProduct.barcode == code,
        ReferenceProduct.id.in_(select(ReferenceProductUnit.reference_product_id).where(ReferenceProductUnit.barcode == code)),
    )))).scalars().first()
    if ref is None:
        return None
    barcodes, names = await company_barcodes(db, company_id)
    return company_item_view(ref, already_owned(ref, barcodes, names))


async def take_for_company(db: AsyncSession, company_id, ids: list) -> dict:
    """The chosen reference products become the company's articles - the same rules as at its creation (its VAT, stock
    0, without price off the Caixa); the ones it has already are skipped."""
    if not ids:
        raise ReferenceCatalogError("Escolha pelo menos um produto do catálogo")
    refs = (await db.execute(select(ReferenceProduct).where(ReferenceProduct.id.in_(ids), ReferenceProduct.is_active.is_(True))
                             .order_by(ReferenceProduct.name))).scalars().all()
    counts, created = await take_reference_products(db, await _company(db, company_id), refs)
    await db.commit()
    return {**counts, "ignorados": len(refs) - counts["artigos"], "product_ids": [str(i) for i in created]}
